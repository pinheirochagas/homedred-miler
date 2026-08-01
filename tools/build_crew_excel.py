#!/usr/bin/env python3
"""Build the editable Homedred Miler aid-station logistics workbook."""

from datetime import datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "crew" / "homedred-crew-logistics.xlsx"

INK = "11110F"
PAPER = "F5F5F1"
WHITE = "FFFFFF"
MUTED = "6D6D67"
LINE = "CECEC5"
LINE_LIGHT = "E5E5DE"
YELLOW = "D7D85B"
YELLOW_SOFT = "F2F2C9"
BLUE = "336F96"
BLUE_SOFT = "DCE8EF"
GREEN = "6D8B65"
GREEN_SOFT = "E2EADF"
GRAY_SOFT = "ECECE7"


STATIONS = [
    {
        "day": "Sat",
        "eta": time(12, 0),
        "mile": 0.0,
        "station": "Start — Golden Gate Park",
        "pacing": "Start group; no dedicated pacer for segments 1–2",
        "to_start": "—",
        "from_finish": "—",
        "supply_driver": "Jonathan",
        "supply_vehicle": "Jonathan's car; begins supply relay",
        "confirm": "Not needed",
        "notes": "Day supplies loaded and labeled.",
    },
    {
        "day": "Sat",
        "eta": time(12, 50),
        "mile": 4.1,
        "station": "GGB south · Welcome Center",
        "pacing": "No pacer handoff",
        "to_start": "—",
        "from_finish": "—",
        "supply_driver": "Jonathan",
        "supply_vehicle": "Jonathan's car",
        "confirm": "Not needed",
        "notes": "",
    },
    {
        "day": "Sat",
        "eta": time(13, 11),
        "mile": 5.9,
        "station": "GGB north · Vista Point",
        "pacing": "Jim starts segments 3–4 to Muir Beach",
        "to_start": "Margarita brings Jim",
        "from_finish": "Margarita picks Jim up at Muir Beach → home/base",
        "supply_driver": "Jonathan",
        "supply_vehicle": "Jonathan's car",
        "confirm": "Unconfirmed",
        "notes": "Jim's CSV says segment 1; meeting says segments 3–4.",
    },
    {
        "day": "Sat",
        "eta": time(14, 18),
        "mile": 10.8,
        "station": "Rodeo Beach",
        "pacing": "Jim continues",
        "to_start": "—",
        "from_finish": "—",
        "supply_driver": "Jonathan",
        "supply_vehicle": "Jonathan's car",
        "confirm": "Unconfirmed",
        "notes": "",
    },
    {
        "day": "Sat",
        "eta": time(16, 6),
        "mile": 17.2,
        "station": "Muir Beach",
        "pacing": "Jim out; Trill starts segments 5–6 to Stinson",
        "to_start": "Margarita brings Trill",
        "from_finish": "Margarita takes Jim home/base; later picks Trill up at Stinson",
        "supply_driver": "Jonathan → Sara",
        "supply_vehicle": "Transfer supplies to Sara's car",
        "confirm": "Unconfirmed",
        "notes": "Jonathan then parks his car at Stinson.",
    },
    {
        "day": "Sat",
        "eta": time(17, 41),
        "mile": 22.9,
        "station": "Pantoll Ranger Station",
        "pacing": "Trill continues",
        "to_start": "—",
        "from_finish": "—",
        "supply_driver": "Sara",
        "supply_vehicle": "Sara's car",
        "confirm": "Unconfirmed",
        "notes": "",
    },
    {
        "day": "Sat",
        "eta": time(18, 27),
        "mile": 26.9,
        "station": "Stinson Beach",
        "pacing": "Trill out; Jonathan starts segment 7 to Bolinas Ridge",
        "to_start": "Jonathan self-drives and parks here",
        "from_finish": "Margarita takes Trill home/base; Max returns Jonathan here",
        "supply_driver": "Sara → Doug",
        "supply_vehicle": "Transfer supplies to Doug's master car",
        "confirm": "Unconfirmed",
        "notes": "Jonathan must extend his stated running window through 22:30.",
    },
    {
        "day": "Sat",
        "eta": time(22, 30),
        "mile": 43.5,
        "station": "Bolinas Ridge Trailhead",
        "pacing": "Jonathan out; Andrey starts segments 8–9 to Sky Oaks",
        "to_start": "Max brings Andrey",
        "from_finish": "Max returns Jonathan to Stinson; Margarita takes Andrey home from Sky Oaks",
        "supply_driver": "Doug",
        "supply_vehicle": "Doug's master car",
        "confirm": "Unconfirmed",
        "notes": "Max must confirm both outbound and return drives.",
    },
    {
        "day": "Sat",
        "eta": time(23, 38),
        "mile": 49.5,
        "station": "Leo T. Cronin roadside handoff",
        "pacing": "Andrey continues",
        "to_start": "—",
        "from_finish": "—",
        "supply_driver": "Doug",
        "supply_vehicle": "Doug's master car; quick roadside handoff",
        "confirm": "Unconfirmed",
        "notes": "No lingering or parking.",
    },
    {
        "day": "Sun",
        "eta": time(3, 45),
        "mile": 65.0,
        "station": "Bolinas Rd × Sky Oaks Rd",
        "pacing": "Andrey out; Doug starts segment 10 to East Peak",
        "to_start": "Doug self-drives master car here",
        "from_finish": "Margarita takes Andrey home/base; Sara retrieves Doug at East Peak",
        "supply_driver": "Doug → Sara",
        "supply_vehicle": "Sara takes Doug's keys and master car",
        "confirm": "Unconfirmed",
        "notes": "Margarita brings Sara here in Margarita's car.",
    },
    {
        "day": "Sun",
        "eta": time(6, 47),
        "mile": 76.1,
        "station": "Mt Tam East Peak",
        "pacing": "Doug out; Pedro continues segment 11 solo",
        "to_start": "—",
        "from_finish": "Sara picks Doug up after the road opens",
        "supply_driver": "Sara",
        "supply_vehicle": "Pickup only after 07:00; aid is not guaranteed",
        "confirm": "Unconfirmed",
        "notes": "Doug waits with phone and warm layer.",
    },
    {
        "day": "Sun",
        "eta": time(7, 39),
        "mile": 80.5,
        "station": "Mountain Home Inn",
        "pacing": "Brian starts segment 12 to Tennessee Valley",
        "to_start": "Jonathan brings Brian",
        "from_finish": "Jonathan picks Brian up at Tennessee Valley → home/base",
        "supply_driver": "Sara",
        "supply_vehicle": "Doug's master car",
        "confirm": "Unconfirmed",
        "notes": "Brian's CSV does not include an availability window.",
    },
    {
        "day": "Sun",
        "eta": time(9, 31),
        "mile": 88.9,
        "station": "Tennessee Valley",
        "pacing": "Brian out; Trill starts segment 13 to GGB north",
        "to_start": "Jonathan brings Trill",
        "from_finish": "Jonathan takes Brian home/base; Sara picks Trill up at GGB north",
        "supply_driver": "Sara",
        "supply_vehicle": "Doug's master car",
        "confirm": "Unconfirmed",
        "notes": "",
    },
    {
        "day": "Sun",
        "eta": time(10, 48),
        "mile": 94.1,
        "station": "GGB north · return",
        "pacing": "Trill out; optional finish group starts",
        "to_start": "—",
        "from_finish": "Sara picks Trill up",
        "supply_driver": "Sara",
        "supply_vehicle": "Doug's master car",
        "confirm": "Unconfirmed",
        "notes": "",
    },
    {
        "day": "Sun",
        "eta": time(11, 8),
        "mile": 95.9,
        "station": "GGB south · return",
        "pacing": "Finish group continues",
        "to_start": "—",
        "from_finish": "—",
        "supply_driver": "Sara",
        "supply_vehicle": "Doug's master car",
        "confirm": "Not needed",
        "notes": "Final bridge-side aid.",
    },
    {
        "day": "Sun",
        "eta": time(12, 0),
        "mile": 100.0,
        "station": "Finish — Golden Gate Park",
        "pacing": "Pedro and finish group arrive",
        "to_start": "—",
        "from_finish": "—",
        "supply_driver": "Sara",
        "supply_vehicle": "Doug's master car; finish and recovery supplies",
        "confirm": "Not needed",
        "notes": "No vehicle should remain at an aid station.",
    },
]

LOCAL_LOGISTICS = {
    "Start — Golden Gate Park": {
        "pacer_in": "—",
        "pacer_out": "Start group",
        "transport": "—",
        "supplies": "Jonathan · day supply car begins",
    },
    "GGB south · Welcome Center": {
        "pacer_in": "—",
        "pacer_out": "—",
        "transport": "—",
        "supplies": "Jonathan · day supply car",
    },
    "GGB north · Vista Point": {
        "pacer_in": "—",
        "pacer_out": "Jim",
        "transport": "Margarita drops Jim",
        "supplies": "Jonathan · day supply car",
    },
    "Rodeo Beach": {
        "pacer_in": "Jim",
        "pacer_out": "Jim",
        "transport": "—",
        "supplies": "Jonathan · day supply car",
    },
    "Muir Beach": {
        "pacer_in": "Jim",
        "pacer_out": "Trill",
        "transport": "Margarita drops Trill and leaves with Jim → home/base",
        "supplies": "Jonathan → Sara · transfer to Sara's car",
    },
    "Pantoll Ranger Station": {
        "pacer_in": "Trill",
        "pacer_out": "Trill",
        "transport": "—",
        "supplies": "Sara · supply car",
    },
    "Stinson Beach": {
        "pacer_in": "Trill",
        "pacer_out": "Jonathan",
        "transport": "Margarita leaves with Trill; Jonathan's car stays here for Max's return",
        "supplies": "Sara → Doug · transfer to master car",
    },
    "Bolinas Ridge Trailhead": {
        "pacer_in": "Jonathan",
        "pacer_out": "Andrey",
        "transport": "Max arrives with Andrey and leaves with Jonathan → Stinson",
        "supplies": "Doug · master supply car",
    },
    "Leo T. Cronin roadside handoff": {
        "pacer_in": "Andrey",
        "pacer_out": "Andrey",
        "transport": "—",
        "supplies": "Doug · quick roadside handoff",
    },
    "Bolinas Rd × Sky Oaks Rd": {
        "pacer_in": "Andrey",
        "pacer_out": "Doug",
        "transport": "Margarita leaves with Andrey → home/base; Sara leaves in Doug's car",
        "supplies": "Doug → Sara · master-car key handoff",
    },
    "Mt Tam East Peak": {
        "pacer_in": "Doug",
        "pacer_out": "— · Pedro continues solo",
        "transport": "Sara leaves with Doug after the road opens",
        "supplies": "No guaranteed aid before 07:00",
    },
    "Mountain Home Inn": {
        "pacer_in": "—",
        "pacer_out": "Brian",
        "transport": "Jonathan drops Brian and drives to Tennessee Valley",
        "supplies": "Sara · master supply car",
    },
    "Tennessee Valley": {
        "pacer_in": "Brian",
        "pacer_out": "Trill",
        "transport": "Jonathan arrives with Trill and leaves with Brian → home/base",
        "supplies": "Sara · master supply car",
    },
    "GGB north · return": {
        "pacer_in": "Trill",
        "pacer_out": "Optional finish group",
        "transport": "Sara leaves with Trill",
        "supplies": "Sara · master supply car",
    },
    "GGB south · return": {
        "pacer_in": "Finish group",
        "pacer_out": "Finish group",
        "transport": "—",
        "supplies": "Sara · final aid",
    },
    "Finish — Golden Gate Park": {
        "pacer_in": "Finish group",
        "pacer_out": "—",
        "transport": "All vehicles complete",
        "supplies": "Sara · finish and recovery supplies",
    },
}


HEADERS = [
    "STOP",
    "DAY",
    "ETA",
    "MILE",
    "AID STATION",
    "PACER IN",
    "PACER OUT",
    "TRANSPORT AT THIS STATION",
    "SUPPLIES AT THIS STATION",
    "CONFIRM",
    "NOTES",
]


def build_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Aid stations"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "F7"
    sheet.sheet_properties.tabColor = YELLOW
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    sheet.merge_cells("A1:K1")
    sheet["A1"] = "HOMEDRED / CREW LOGISTICS"
    sheet["A1"].font = Font(name="Arial Narrow", size=26, color=INK)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 40

    sheet.merge_cells("A2:K2")
    sheet["A2"] = "Aid-station assignments · editable start · fixed 24-hour target"
    sheet["A2"].font = Font(name="Arial", size=11, color=MUTED)
    sheet["A2"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 23

    sheet["A3"] = "START"
    sheet["A3"].font = Font(name="Arial", size=9, bold=True, color=WHITE)
    sheet["A3"].fill = PatternFill("solid", fgColor=INK)
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells("B3:D3")
    sheet["B3"] = datetime(2026, 8, 8, 12, 0)
    sheet["B3"].number_format = 'ddd, mmm d yyyy "·" hh:mm'
    sheet["B3"].font = Font(name="Arial", size=11, bold=True, color=INK)
    sheet["B3"].fill = PatternFill("solid", fgColor=YELLOW)
    sheet["B3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["B3"].comment = Comment(
        "Edit this date and time. Every ETA and day below updates automatically.",
        "Homedred",
    )

    sheet.merge_cells("E3:K3")
    sheet["E3"] = "Edit the yellow start field; ETA formulas update automatically."
    sheet["E3"].font = Font(name="Arial", size=9, color=MUTED)
    sheet["E3"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[3].height = 29

    sheet.merge_cells("A4:K4")
    sheet["A4"] = (
        "ONE ROW = ONE AID-STATION VISIT. Every cell describes only what happens "
        "at that station. Edit cells or insert/delete rows as needed."
    )
    sheet["A4"].font = Font(name="Arial", size=9, color=INK)
    sheet["A4"].fill = PatternFill("solid", fgColor=YELLOW_SOFT)
    sheet["A4"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[4].height = 25

    header_row = 6
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    sheet.cell(header_row, 8).fill = PatternFill("solid", fgColor=GREEN)
    sheet.cell(header_row, 9).fill = PatternFill("solid", fgColor=BLUE)
    sheet.cell(header_row, 10).fill = PatternFill("solid", fgColor=YELLOW)
    sheet.cell(header_row, 10).font = Font(name="Arial", size=9, bold=True, color=INK)
    sheet.row_dimensions[header_row].height = 34

    comments = {
        6: "Pacer arriving at this station with Pedro.",
        7: "Pacer leaving this station with Pedro.",
        8: "Only the vehicle movement that happens at this station.",
        9: "Only the supply-car action that happens at this station.",
        10: "Choose Unconfirmed, Confirmed, or Not needed.",
    }
    for column, text in comments.items():
        sheet.cell(header_row, column).comment = Comment(text, "Homedred")

    thin = Side(style="thin", color=LINE)
    light = Side(style="thin", color=LINE_LIGHT)

    for index, station in enumerate(STATIONS, start=1):
        row = header_row + index
        local = LOCAL_LOGISTICS[station["station"]]
        day_offset = 0 if station["day"] == "Sat" else 24 * 60
        eta = station["eta"]
        offset_minutes = day_offset + eta.hour * 60 + eta.minute - 12 * 60
        values = [
            index,
            f'=TEXT(C{row},"ddd")',
            f"=$B$3+{offset_minutes}/1440",
            station["mile"],
            station["station"],
            local["pacer_in"],
            local["pacer_out"],
            local["transport"],
            local["supplies"],
            station["confirm"],
            station["notes"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=light)
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=WHITE)

        sheet.cell(row, 1).font = Font(name="Arial Narrow", size=12, color=MUTED)
        sheet.cell(row, 2).font = Font(name="Arial", size=9, bold=True, color=MUTED)
        sheet.cell(row, 3).number_format = "hh:mm"
        sheet.cell(row, 3).font = Font(name="Arial Narrow", size=12, color=INK)
        sheet.cell(row, 4).number_format = '0.0" mi"'
        sheet.cell(row, 4).font = Font(name="Arial", size=9, color=MUTED)
        sheet.cell(row, 5).font = Font(name="Arial", size=10, bold=True, color=INK)
        sheet.cell(row, 10).font = Font(name="Arial", size=9, bold=True, color=INK)
        sheet.row_dimensions[row].height = 54

    last_row = header_row + len(STATIONS)
    sheet.auto_filter.ref = f"A{header_row}:K{last_row}"

    status_list = workbook.create_sheet("_lists")
    status_list.sheet_state = "hidden"
    for row, status in enumerate(("Unconfirmed", "Confirmed", "Not needed"), start=1):
        status_list.cell(row, 1, status)

    validation = DataValidation(
        type="list",
        formula1="'_lists'!$A$1:$A$3",
        allow_blank=False,
    )
    validation.error = "Choose a confirmation status from the list."
    validation.errorTitle = "Invalid status"
    validation.prompt = "Select Unconfirmed, Confirmed, or Not needed."
    validation.promptTitle = "Confirmation status"
    sheet.add_data_validation(validation)
    validation.add("J7:J500")

    sheet.conditional_formatting.add(
        "J7:J500",
        CellIsRule(
            operator="equal",
            formula=['"Unconfirmed"'],
            fill=PatternFill("solid", fgColor=YELLOW_SOFT),
        ),
    )
    sheet.conditional_formatting.add(
        "J7:J500",
        CellIsRule(
            operator="equal",
            formula=['"Confirmed"'],
            fill=PatternFill("solid", fgColor=GREEN_SOFT),
        ),
    )
    sheet.conditional_formatting.add(
        "J7:J500",
        CellIsRule(
            operator="equal",
            formula=['"Not needed"'],
            fill=PatternFill("solid", fgColor=GRAY_SOFT),
        ),
    )

    widths = {
        "A": 7,
        "B": 7,
        "C": 9,
        "D": 9,
        "E": 27,
        "F": 18,
        "G": 24,
        "H": 48,
        "I": 36,
        "J": 16,
        "K": 38,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for cell in sheet[header_row]:
        cell.border = Border(top=thin, bottom=thin)

    sheet.print_area = f"A1:K{last_row}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4
    sheet.oddFooter.center.text = "HOMEDRED / CREW LOGISTICS"
    sheet.oddFooter.right.text = "Page &P of &N"

    workbook.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=False)
    main = check["Aid stations"]
    assert main.max_row == last_row
    assert main.max_column == len(HEADERS)
    assert main["B3"].value == datetime(2026, 8, 8, 12, 0)
    assert main["E7"].value == "Start — Golden Gate Park"
    assert main["C7"].value == "=$B$3+0/1440"
    assert main[f"E{last_row}"].value == "Finish — Golden Gate Park"
    assert len(main.data_validations.dataValidation) == 1
    check.close()

    print(OUTPUT)


if __name__ == "__main__":
    build_workbook()
